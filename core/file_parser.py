"""
Đọc nội dung đề thi từ file .docx hoặc .pdf (đúng 2 định dạng
Ban Tổ chức công bố sẽ dùng theo Điều 3 Thể lệ cuộc thi).
"""

import pathlib
from docx import Document
import pdfplumber


def read_docx(path: str) -> str:
    doc = Document(path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def read_pdf(path: str) -> str:
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)


def read_input_file(path: str) -> str:
    ext = pathlib.Path(path).suffix.lower()
    if ext == ".docx":
        text = read_docx(path)
    elif ext == ".pdf":
        text = read_pdf(path)
    else:
        raise ValueError(f"Định dạng file không được hỗ trợ: {ext} (chỉ nhận .docx hoặc .pdf)")

    if not text.strip():
        raise ValueError(
            "Không trích xuất được nội dung văn bản từ file. "
            "Có thể đây là file scan/ảnh — cần OCR trước khi đưa vào hệ thống."
        )
    return text

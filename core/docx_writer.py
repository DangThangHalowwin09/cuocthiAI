"""
Xuất văn bản kết quả (Cáo trạng / Phát biểu của Kiểm sát viên) ra file
.docx ĐÚNG THỂ THỨC theo Phụ lục I, Nghị định số 30/2020/NĐ-CP ngày
05/3/2020 của Chính phủ về công tác văn thư.
"""

import os
from pathlib import Path

from docx import Document
from docx.shared import Pt, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "Times New Roman"

TITLES = {
    "HINH_SU": "CÁO TRẠNG",
    "DAN_SU_HANH_CHINH": "PHÁT BIỂU",
}

SUBTITLES = {
    "HINH_SU": None,
    "DAN_SU": "Của Kiểm sát viên tại phiên tòa (phiên họp) sơ thẩm",
    "HANH_CHINH": "Của Kiểm sát viên tại phiên tòa hành chính sơ thẩm",
}

DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / "data" / "templates" / "156-Cáo trạng.docx"
)


def _set_font(run, size: int, bold: bool = False, italic: bool = False):
    run.font.name = FONT_NAME
    # Đảm bảo áp dụng đúng font cho cả ký tự Đông Á
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _add_run(paragraph, text: str, size: int, bold: bool = False, italic: bool = False):
    run = paragraph.add_run(text)
    _set_font(run, size, bold, italic)
    return run


def _set_single_spacing(paragraph, space_after: int = 0, space_before: int = 0, indent: bool = False):
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent:
        pf.first_line_indent = Cm(1.27)  # Thụt lề đầu dòng chuẩn 1.27 cm (0.5 inch)


def _add_bottom_border(paragraph):
    """Thêm đường kẻ ngang bên dưới đoạn văn."""
    p_el = paragraph._p
    pPr = p_el.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_margins(doc: Document):
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)


def _iter_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_template_token(doc: Document, token: str, value: str) -> bool:
    replaced = False
    for paragraph in _iter_paragraphs(doc):
        if token in paragraph.text:
            paragraph.text = paragraph.text.replace(token, value)
            replaced = True
    return replaced


def _add_body_paragraph_after(marker: Paragraph, text: str) -> Paragraph:
    paragraph_element = OxmlElement("w:p")
    marker._p.addnext(paragraph_element)
    paragraph = Paragraph(paragraph_element, marker._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    is_can_cu = text.startswith("Căn cứ")
    is_section_header = (
        text.upper() == text and len(text) < 60 and any(c.isalpha() for c in text)
    )

    if is_section_header:
        _set_single_spacing(paragraph, space_before=6, space_after=6)
        _add_run(paragraph, text, size=14, bold=True)
    else:
        _set_single_spacing(paragraph, space_after=6, indent=True)
        _add_run(paragraph, text, size=14, italic=is_can_cu)
    return paragraph


def _remove_paragraph(paragraph: Paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def _replace_sample_body(doc: Document, final_text: str) -> bool:
    paragraphs = list(doc.paragraphs)
    start_index = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip().startswith("Căn cứ các điều")
        ),
        None,
    )
    end_index = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip().startswith("……………………")
        ),
        None,
    )
    if start_index is None or end_index is None or start_index >= end_index:
        return False

    anchor = paragraphs[start_index - 1]
    for paragraph in paragraphs[start_index : end_index + 1]:
        _remove_paragraph(paragraph)

    previous = anchor
    for line in final_text.splitlines():
        line = line.strip()
        if line:
            previous = _add_body_paragraph_after(previous, line)
    return True


def _write_from_template(
    final_text: str,
    output_path: str,
    template_path: str,
    unit_name: str,
    unit_parent: str,
) -> str:
    template = Path(template_path)
    if template.suffix.lower() != ".docx":
        raise ValueError(
            "Mẫu 156 phải được chuyển sang .docx trước khi dùng làm template. "
            f"Định dạng nhận được: {template.suffix or '(không có phần mở rộng)'}"
        )
    if not template.is_file():
        raise FileNotFoundError(f"Không tìm thấy template mẫu 156: {template}")

    doc = Document(str(template))
    if not _replace_sample_body(doc, final_text):
        token_values = {
            "{{DON_VI_BAN_HANH}}": unit_name,
            "{{DON_VI_CAP_TREN}}": unit_parent,
            "{{SO_KY_HIEU}}": "Số: ..../CT-VKS...",
            "{{DIA_DANH_NGAY}}": "Nghệ An, ngày .... tháng .... năm 20....",
            "{{CHUC_DANH_NGUOI_KY}}": "VIỆN TRƯỞNG",
        }
        for token, value in token_values.items():
            if not _replace_template_token(doc, token, value):
                raise ValueError(f"Template mẫu 156 thiếu vùng nội dung: {token}")
        marker = next(
            (
                paragraph
                for paragraph in _iter_paragraphs(doc)
                if "{{NOI_DUNG_CAO_TRANG}}" in paragraph.text
            ),
            None,
        )
        if marker is None:
            raise ValueError(
                "Template mẫu 156 thiếu vùng nội dung: {{NOI_DUNG_CAO_TRANG}}"
            )
        marker.text = ""
        previous = marker
        for line in final_text.splitlines():
            line = line.strip()
            if line:
                previous = _add_body_paragraph_after(previous, line)

    doc.save(output_path)
    return output_path


def _add_header_table(doc: Document, unit_name: str, unit_parent: str, so_ky_hieu: str, dia_danh_ngay: str):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left_cell, right_cell = table.rows[0].cells

    table.style = None

    # --- CỘT TRÁI: tên cơ quan chủ quản + tên cơ quan ban hành + số ký hiệu ---
    p1 = left_cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p1)
    _add_run(p1, unit_parent, size=12, bold=False)

    p2 = left_cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p2)
    _add_run(p2, unit_name, size=12, bold=True)
    _add_bottom_border(p2)

    p3 = left_cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p3)
    _add_run(p3, so_ky_hieu, size=13, bold=False)

    # --- CỘT PHẢI: Quốc hiệu, Tiêu ngữ, Địa danh ngày tháng ---
    q1 = right_cell.paragraphs[0]
    q1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(q1)
    _add_run(q1, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", size=12, bold=True)

    q2 = right_cell.add_paragraph()
    q2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(q2)
    _add_run(q2, "Độc lập - Tự do - Hạnh phúc", size=13, bold=True)
    _add_bottom_border(q2)

    q3 = right_cell.add_paragraph()
    q3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(q3)
    _add_run(q3, dia_danh_ngay, size=13, italic=True)

    return table


def write_result_docx(
    final_text: str,
    case_type: str,
    output_path: str,
    is_hanh_chinh: bool = False,
    unit_name: str = "VIỆN KIỂM SÁT NHÂN DÂN TỈNH NGHỆ AN",
    unit_parent: str = "VIỆN KIỂM SÁT NHÂN DÂN TỐI CAO",
    template_path: str | None = None,
) -> str:
    template_path = template_path or os.environ.get("CAO_TRANG_TEMPLATE_PATH")
    if template_path or DEFAULT_TEMPLATE_PATH.is_file():
        return _write_from_template(
            final_text,
            output_path,
            template_path or str(DEFAULT_TEMPLATE_PATH),
            unit_name,
            unit_parent,
        )
    if case_type == "HINH_SU":
        raise FileNotFoundError(
            "Chưa cấu hình template mẫu 156. Đặt file DOCX tại "
            f"{DEFAULT_TEMPLATE_PATH} hoặc cấu hình CAO_TRANG_TEMPLATE_PATH."
        )

    doc = Document()
    _set_margins(doc)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = FONT_NAME
    normal_style.font.size = Pt(14)

    ky_hieu = "CT" if case_type == "HINH_SU" else ("PB-HC" if is_hanh_chinh else "PB")
    so_ky_hieu = f"Số: ..../{ky_hieu}-VKS..."
    dia_danh_ngay = "Nghệ An, ngày .... tháng .... năm 20...."

    _add_header_table(doc, unit_name, unit_parent, so_ky_hieu, dia_danh_ngay)

    doc.add_paragraph()  # dòng trống

    # --- TÊN LOẠI VĂN BẢN ---
    title_text = TITLES.get(case_type, "VĂN BẢN")
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p_title)
    _add_run(p_title, title_text, size=14, bold=True)

    # --- TRÍCH YẾU (nếu có) ---
    subtitle_key = "HANH_CHINH" if is_hanh_chinh else ("DAN_SU" if case_type != "HINH_SU" else None)
    subtitle = SUBTITLES.get(subtitle_key) if subtitle_key else None
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_single_spacing(p_sub)
        _add_run(p_sub, subtitle, size=14, bold=True)

    doc.add_paragraph()  # dòng trống trước nội dung

    # --- NỘI DUNG VĂN BẢN ---
    for line in final_text.split("\n"):
        line = line.strip()
        if not line:
            continue

        p = doc.add_paragraph()
        is_can_cu = line.startswith("Căn cứ")
        is_section_header = (
            line.upper() == line and len(line) < 60 and any(c.isalpha() for c in line)
        )

        if is_section_header:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_single_spacing(p, space_before=6, space_after=6, indent=False)
            _add_run(p, line, size=14, bold=True)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # Căn đều 2 bên chuẩn văn bản hành chính
            _set_single_spacing(p, space_after=6, indent=True)  # Thụt lề đầu dòng
            _add_run(p, line, size=14, italic=is_can_cu)

    doc.add_paragraph()

    # --- CHỨC VỤ, HỌ TÊN NGƯỜI KÝ + NƠI NHẬN ---
    footer_table = doc.add_table(rows=1, cols=2)
    footer_table.autofit = True
    nhan_cell, ky_cell = footer_table.rows[0].cells

    p_nn_title = nhan_cell.paragraphs[0]
    _set_single_spacing(p_nn_title)
    _add_run(p_nn_title, "Nơi nhận:", size=12, bold=True, italic=True)

    for line in ["- Như trên;", "- Lãnh đạo đơn vị (để báo cáo);", "- Lưu: VT, HSKS."]:
        p_nn = nhan_cell.add_paragraph()
        _set_single_spacing(p_nn)
        _add_run(p_nn, line, size=11)

    p_ky_title = ky_cell.paragraphs[0]
    p_ky_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p_ky_title)
    chuc_danh = "VIỆN TRƯỞNG" if case_type == "HINH_SU" else "KIỂM SÁT VIÊN"
    _add_run(p_ky_title, chuc_danh, size=14, bold=True)

    for _ in range(3):
        ky_cell.add_paragraph()

    p_ten = ky_cell.add_paragraph()
    p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_single_spacing(p_ten)
    _add_run(p_ten, "(Ký, ghi rõ họ tên, đóng dấu)", size=13, italic=True)

    doc.add_paragraph()
    note = doc.add_paragraph()
    _add_run(
        note,
        "* Văn bản do AI hỗ trợ soạn thảo theo thể thức Nghị định 30/2020/NĐ-CP. "
        "Kiểm sát viên có trách nhiệm kiểm tra, chỉnh sửa trước khi ban hành chính thức.",
        size=9,
        italic=True,
    )

    doc.save(output_path)
    return output_path